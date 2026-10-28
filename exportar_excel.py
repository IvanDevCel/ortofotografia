"""Public Python dependency; no private Codex/Node artifact-tool required."""
import json,argparse
from pathlib import Path
def main():
    p=argparse.ArgumentParser();p.add_argument('informe',type=Path);args=p.parse_args()
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    from openpyxl.utils import get_column_letter
    out=args.informe;wb=Workbook();wb.remove(wb.active)
    for filename,title in [('resumen.json','Resumen'),('filas.json','Vuelo_1Hz')]:
        data=json.loads((out/filename).read_text(encoding='utf8'));sh=wb.create_sheet(title)
        if isinstance(data,dict):rows=[['Medida','Valor']]+[[k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v] for k,v in data.items()]
        else:
            keys=list(data[0]);rows=[keys]+[[r.get(k) for k in keys] for r in data]
        for row in rows:sh.append(row)
        for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='293D50')
        sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
        for i in range(1,sh.max_column+1):sh.column_dimensions[get_column_letter(i)].width=27 if title!='Resumen' else 65
    wb.save(out/'informe.xlsx');print(out/'informe.xlsx')
if __name__=='__main__':main()
